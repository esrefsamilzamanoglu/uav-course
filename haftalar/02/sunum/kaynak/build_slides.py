"""Build editable 16:9 week-2 slides from validated Gazebo evidence.
Requires: python-pptx, Pillow. Run prepare_evidence.py only when revalidating.
"""
from pathlib import Path
import json,math,re
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from PIL import Image
from curriculum import SLIDES,SOURCES
from revise_curriculum import revise
SLIDES=revise(SLIDES,SOURCES)
def typography(v):
    if isinstance(v,str):
        v=re.sub(r'(?<=[0-9])(?=[a-zçğıöşüA-CE-Z])',' ',v)
        v=re.sub(r'(?<=[a-zçğıöşü])(?=[23]D)',' ',v)
        return re.sub(r'(?<=[23]D)(?=[a-zçğıöşü])',' ',v)
    if isinstance(v,dict):return {k:(val if k in ('image','images','sources','code','diagram','layout','section') else typography(val)) for k,val in v.items()}
    if isinstance(v,(tuple,list)):return [typography(x) for x in v]
    return v
SLIDES=typography(SLIDES)
for lesson in SLIDES:
    if lesson.get('layout')=='code':
        # Explicit PowerShell continuation keeps the displayed multi-line command runnable.
        lesson['code']=lesson['code'].replace('\n  ',' ').replace(' .\\tools',' `\n  .\\tools').replace(' "UCUS_KLASORU"',' `\n  "UCUS_KLASORU"')
ROOT=Path(__file__).resolve().parents[2];ASSETS=ROOT/'sunum/gorseller'
GENERATED=ROOT/'sunum/uretilen_gorseller'
ILLUSTRATIONS=json.loads((GENERATED/'prompts.json').read_text(encoding='utf-8'))
ILLUSTRATED={a['slide']:a for a in ILLUSTRATIONS['assets']}
assert json.loads((ROOT/'dogrulama/final_summary.json').read_text(encoding='utf-8'))['all_required_passed'], 'Validate simulations before making slides'
assert 1<len(SLIDES)<=100
P=Presentation();P.slide_width=Inches(13.333333);P.slide_height=Inches(7.5)
P.core_properties.title='İHA Yazılımı ve Simülasyonu — 2. Hafta: AUTO Görevler'
P.core_properties.author='Arş. Gör. Eşref Şamil Zamanoğlu'
P.core_properties.subject='Birlik Akademi · Bilişim Programları · 2026 Güz Dönemi'
BG='EDF3FF';NAVY='18365D';TEAL='238B88';GOLD='D79A29';INK='263D56';MUTED='61758C';WHITE='FFFFFF';PALE='DCE8F7';RED='B74A47'
def rgb(v):return RGBColor.from_string(v)
def rect(s,x,y,w,h,fill=WHITE,line=None,radius=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if line:sh.line.color.rgb=rgb(line)
    else:sh.line.fill.background()
    if radius:
        try:sh.adjustments[0]=.09
        except Exception:pass
    return sh
def text(s,value,x,y,w,h,size=20,color=INK,bold=False,align=None,font='Century Gothic'):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.03);tf.margin_top=Inches(.02);tf.margin_bottom=0
    for i,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(6)
        if align is not None:p.alignment=align
    return sh
def image(s,name,x,y,w,h,fill=False):
    path=ASSETS/name
    if fill:s.shapes.add_picture(str(path),Inches(x),Inches(y),Inches(w),Inches(h));return
    iw,ih=Image.open(path).size;k=min(w/iw,h/ih);ww,hh=iw*k,ih*k
    s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),Inches(ww),Inches(hh))
def line(s,x1,y1,x2,y2,color=TEAL,width=2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width);return sh
def circle(s,x,y,d,fill=TEAL):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.fill.background();return sh
def arrow(s,x,y,w=.4,h=.3):
    a=s.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x),Inches(y),Inches(w),Inches(h));a.fill.solid();a.fill.fore_color.rgb=rgb(TEAL);a.line.fill.background()
def card(s,title,body,x,y,w,h=1.22,n=None):
    rect(s,x,y,w,h,WHITE,radius=True)
    text(s,title,x+.18,y+.13,w-.36,.45,20,NAVY,True)
    text(s,body,x+.18,y+.64,w-.36,h-.67,17,INK)
def base(d,n):
    s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
    rect(s,.57,.42,.1,.56,TEAL);text(s,d['title'],.85,.37,11.8,.8,29,NAVY,True)
    text(s,d.get('section','AUTO UYGULAMALARI').upper(),.61,1.12,12,.25,10,MUTED,True)
    rect(s,.6,6.55,12.1,.5,PALE,radius=True);text(s,d['message'],.78,6.64,11.75,.32,16,NAVY,True)
    text(s,'Birlik Akademi  ·  İHA Yazılımı ve Simülasyonu  ·  2. Hafta',.62,7.18,10,.2,9,MUTED)
    text(s,f'{n:02d} / {len(SLIDES):02d}',11.85,7.15,.85,.25,10,MUTED,align=PP_ALIGN.RIGHT)
    return s
def draw_table(s,headers,rows):
    x,y,w=.65,1.65,12.02;cols=len(headers);widths=[w/cols]*cols
    rh=min(.76,4.45/(len(rows)+1));pos=x
    for head,cw in zip(headers,widths):rect(s,pos,y,cw-.04,rh,NAVY);text(s,head,pos+.13,y+.13,cw-.26,rh-.12,17,WHITE,True);pos+=cw
    for j,row in enumerate(rows):
        pos=x;yy=y+(j+1)*rh
        for value,cw in zip(row,widths):rect(s,pos,yy,cw-.04,rh-.04,WHITE if j%2==0 else PALE);text(s,value,pos+.13,yy+.1,cw-.26,rh-.13,16,INK);pos+=cw
def diagram(s,kind):
    if kind=='fov':
        rect(s,1,5.7,8,.12,TEAL);circle(s,4.35,1.65,.3,NAVY)
        line(s,4.5,1.95,1.65,5.7);line(s,4.5,1.95,7.35,5.7);line(s,4.5,2.1,4.5,5.7,MUTED,1)
        text(s,'Kamera',4.8,1.65,2,.4,19,NAVY,True);text(s,'70°',4,2.7,1,.5,24,TEAL,True);text(s,'18 m',4.65,4.1,1.5,.4,20)
        text(s,'25,2 m yatay ayak izi',2.4,5.95,5,.4,20,NAVY,True)
    elif kind in ('grid','spray_grid'):
        rect(s,1.2,1.75,7.3,4.2,'D6E9D8');last=None
        count=9 if kind=='spray_grid' else 7
        for i in range(count):
            x=1.7+i*6.24/(count-1);ya,yb=(2.1,5.65) if i%2==0 else (5.65,2.1)
            if last:line(s,*last,x,ya,GOLD,3)
            line(s,x,ya,x,yb,TEAL,3)
            for yy in [2.25,2.9,3.55,4.2,4.85,5.5]:circle(s,x-.04,yy,.08,NAVY)
            last=(x,yb)
        text(s,('7' if kind=='spray_grid' else '8')+' m şerit aralığı',2.4,6,4,.35,18,NAVY,True)
    elif kind=='orbit':
        cx,cy,rx,ry=4.6,3.7,3.1,1.55
        for offset,color in [(0,TEAL),(-.65,GOLD)]:
            prev=None
            for i in range(49):
                a=2*math.pi*i/48;p=(cx+rx*math.cos(a),cy+ry*math.sin(a)+offset)
                if prev:line(s,*prev,*p,color,1.8)
                prev=p
        rect(s,3.95,2.9,1.3,1.7,NAVY);text(s,'BİNA',4.05,3.5,1.1,.35,17,WHITE,True)
        for x,y in [(1.5,3.7),(6.8,2.6),(7.3,4.4)]:circle(s,x-.08,y-.08,.17,GOLD);line(s,x,y,4.6,3.7,MUTED,1)
        text(s,'18 m / 23 m · R = 26 m',2.2,5.8,5,.4,20,NAVY,True)
    elif kind=='pipeline':pass
    elif kind=='coordinates':
        circle(s,3.4,4.8,.22,NAVY);line(s,3.5,4.9,7.7,4.9,TEAL,4);line(s,3.5,4.9,3.5,1.9,TEAL,4);line(s,3.5,4.9,1.7,3.3,GOLD,4)
        text(s,'E / x',7.8,4.65,1,.4,23,TEAL,True);text(s,'N / y',3.1,1.4,1.8,.4,23,TEAL,True);text(s,'U / z',.7,2.95,1.6,.4,23,GOLD,True);text(s,'HOME',3.15,5.25,2,.5,20,NAVY,True)
    elif kind=='landing':
        rect(s,.95,5.35,7.8,.12,NAVY)
        for p in [(1.5,2.2),(4.5,3.45),(7.8,5.35)]:circle(s,p[0]-.08,p[1]-.08,.16,TEAL)
        line(s,1.5,2.2,4.5,3.45,TEAL,3);line(s,4.5,3.45,7.8,5.35,TEAL,3)
        text(s,'60 m',1.1,1.65,1.7,.4,20,NAVY,True);text(s,'30 m',4.1,2.9,1.7,.4,20,NAVY,True);text(s,'LAND',7.1,5.65,1.8,.4,20,NAVY,True)
        text(s,'450 m',1.2,5.65,2,.4,18);text(s,'240 m',4.25,5.65,2,.4,18)

TABLE_WIDTHS={4:[1.95,3,3.2],5:[3.3,4.85],8:[2.45,2.4,3.3],15:[1.35,3.45,3.35],
 16:[.7,3.9,3.55],20:[1.7,3.5,2.95],22:[2.85,3,2.3],25:[2.25,2.9,3],
 27:[1.9,3.25,3],38:[1.9,3.7,2.55],63:[2.4,3.35,2.4],70:[2.05,2.8,3.3],
 79:[1.9,2.6,3.65],86:[2.2,3,2.95],88:[1.9,2.5,3.75],89:[2.6,2.85,2.7],93:[2.85,5.3]}
def teaching_picture(s,n,x,y,w,h):
    asset=ILLUSTRATED[n]
    assert asset['title']==SLIDES[n-1]['title'], 'Illustration/slide mismatch'
    path=(GENERATED/asset['file']).with_suffix('.jpg');assert path.is_file(),path
    iw,ih=Image.open(path).size;k=min(w/iw,(h-.28)/ih);ww,hh=iw*k,ih*k
    s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-.28-hh)/2),Inches(ww),Inches(hh))
    text(s,'Kavramsal çizim · AI',x,y+h-.22,w,.2,9,MUTED,align=PP_ALIGN.CENTER)
def illustrated_table(s,d,n):
    teaching_picture(s,n,9.13,1.55,3.55,4.8)
    widths=TABLE_WIDTHS[n];x=.65;y=1.65;header=.7;rh=3.9/len(d['rows'])
    for head,cw in zip(d['headers'],widths):
        rect(s,x,y,cw-.04,header,NAVY);text(s,head,x+.1,y+.1,cw-.23,header-.12,15,WHITE,True);x+=cw
    for j,row in enumerate(d['rows']):
        x=.65;yy=y+header+j*rh
        for value,cw in zip(row,widths):
            rect(s,x,yy,cw-.04,rh-.04,WHITE if j%2==0 else PALE)
            text(s,value,x+.1,yy+.09,cw-.23,rh-.13,15,INK);x+=cw
def illustrated_cards(s,d,n):
    items=d['cards'];left=n%2==0
    teaching_picture(s,n,.68 if left else 7.75,1.5,4.85,4.9)
    x=5.8 if left else .68;w=6.85
    if len(items)<=3:
        h=1.43
        for j,(title,body) in enumerate(items):
            y=1.65+j*1.55;rect(s,x,y,w,h,WHITE,radius=True)
            circle(s,x+.15,y+.18,.39,TEAL);text(s,str(j+1),x+.24,y+.22,.22,.25,12,WHITE,True)
            text(s,title,x+.7,y+.13,w-.9,.43,20,NAVY,True)
            text(s,body,x+.7,y+.65,w-.9,.7,18)
    else:
        for j,(title,body) in enumerate(items):
            xx=x+(j%2)*3.51;yy=1.7+(j//2)*2.3
            rect(s,xx,yy,3.34,2.13,WHITE,radius=True)
            text(s,title,xx+.17,yy+.16,3,.56,19,NAVY,True)
            text(s,body,xx+.17,yy+.9,3,1.04,17)
def illustrated_flow(s,d,n):
    teaching_picture(s,n,.72,1.48,5.2,4.9)
    items=d['cards'];x=6.2;w=6.4;h=4.55/len(items)
    for j,(title,body) in enumerate(items):
        y=1.65+j*h
        if j<len(items)-1:line(s,x+.27,y+.42,x+.27,y+h+.1,TEAL,2)
        rect(s,x+.73,y,w-.73,h-.12,WHITE,radius=True)
        circle(s,x,y+.14,.53,TEAL);text(s,str(j+1),x+.16,y+.21,.26,.27,15,WHITE,True)
        text(s,title,x+.92,y+.09,w-1.15,.37,19,NAVY,True)
        text(s,body,x+.92,y+.55,w-1.15,h-.68,17)
def illustrated_code(s,d,n):
    teaching_picture(s,n,8.15,1.52,4.48,4.86)
    rect(s,.7,1.75,7.17,2.04,NAVY,radius=True)
    text(s,'PowerShell · simulasyon klasöründen',.91,1.87,6.77,.27,10,'BDCEE2')
    text(s,d['code'],.91,2.23,6.77,1.35,16,WHITE,font='Consolas')
    for i,(title,body) in enumerate(d['cards']):
        y=4.01+i*1.13;rect(s,.7,y,7.17,1.02,WHITE,radius=True)
        text(s,title,.88,y+.1,6.8,.32,18,NAVY,True);text(s,body,.88,y+.49,6.8,.45,16)
def illustrated_formula(s,d,n):
    teaching_picture(s,n,7.65,1.55,4.98,4.84)
    text(s,d['formula'],.8,1.8,6.6,.7,31,NAVY,True)
    text(s,'20 L/ha × 10,8 km/h × 8 m / 600\n= 2,88 L/min',.82,2.75,6.5,1.03,24,TEAL,True)
    for i,(var,label) in enumerate(d['legend']):
        y=4.05+i*.73;circle(s,.88,y,.44,TEAL);text(s,var,1,y+.06,.3,.3,17,WHITE,True);text(s,label,1.53,y+.04,5.75,.5,18)
def illustrated_stats(s,d,n):
    teaching_picture(s,n,.68,1.5,5.25,4.87)
    for j,(value,label) in enumerate(d['stats']):
        y=1.65+j*1.4;rect(s,6.2,y,6.4,1.2,WHITE,radius=True)
        text(s,value,6.4,y+.13,1.45,.8,43,TEAL if j<2 else GOLD,True)
        text(s,label,8.05,y+.3,4.3,.7,19,NAVY,True)
    text(s,'Olay sayısı ≠ kişi sayısı',6.38,6,6.1,.35,17,MUTED)
def render(d,n):
    s=base(d,n);layout=d.get('layout','photo');items=d.get('cards',[])
    if n in ILLUSTRATED:
        {'table':illustrated_table,'cards':illustrated_cards,'flow':illustrated_flow,
         'code':illustrated_code,'formula':illustrated_formula,'stats':illustrated_stats}[layout](s,d,n)
    elif layout=='chapter':
        s.background.fill.fore_color.rgb=rgb(NAVY)
        for sh in list(s.shapes):
            el=sh._element;el.getparent().remove(el)
        text(s,d['chapter'],.7,.45,3.8,1.8,86,'385777',True)
        text(s,d['title'],.75,2.3,4.15,2.1,29,WHITE,True)
        text(s,d.get('subtitle',d['message']),.78,4.8,3.9,1.25,18,'9DDED7')
        image(s,d['image'],5.1,.8,7.6,5.8)
        rect(s,.78,6.7,.65,.08,TEAL)
        text(s,d['message'],1.65,6.58,10.8,.58,17,WHITE)
        text(s,f'{n:02d} / {len(SLIDES):02d}',11.8,7.18,.85,.2,9,'BDCEE2',align=PP_ALIGN.RIGHT)
    elif layout=='cover':
        s.background.fill.fore_color.rgb=rgb(NAVY)
        for sh in list(s.shapes):
            el=sh._element;el.getparent().remove(el)
        text(s,'İHA YAZILIMI\nVE SİMÜLASYONU',.75,.65,8,1.6,36,WHITE,True)
        text(s,'2. HAFTA  /  AUTO GÖREVLER',.8,2.4,8,.55,24,'8DD8D1',True)
        text(s,'Gerçek problemlerden simülasyon uygulamalarına',.8,3.05,11,.5,21,WHITE)
        for i,name in enumerate(['mapping_dunya.png','gaussian_final.png','plane_dunya.png']):image(s,name,.8+i*4.02,3.85,3.85,2.3)
        text(s,'Birlik Akademi · Bilişim Programları  |  2026 Güz Dönemi',.8,6.33,11.8,.4,16,WHITE)
        text(s,'Arş. Gör. Eşref Şamil Zamanoğlu  ·  Birlik Vakfı Sakarya Şubesi',.8,6.85,11.8,.35,15,'C7D9EF')
    elif layout=='visual':
        image(s,d['image'],.62,1.4,12.08,4.98)
    elif layout=='gallery':
        for i,name in enumerate(d['images']):
            x=.65+i*4.12;image(s,name,x,1.6,3.8,3.9)
            title,body=items[i];text(s,title,x+.03,5.6,3.75,.42,20,NAVY,True);text(s,body,x+.03,6.05,3.75,.35,15)
    elif layout=='spotlight':
        image(s,d['image'],.7,1.5,8.15,4.85)
        for i,(value,label) in enumerate(d['stats']):
            y=1.65+i*1.62;text(s,value,9.15,y,3.4,.8,40,TEAL,True);text(s,label,9.18,y+.9,3.35,.45,16,NAVY)
    elif layout=='stats':
        for i,(value,label) in enumerate(d['stats']):
            x=.7+i*4.15;text(s,value,x,2.1,3.8,1.6,76,TEAL if i<2 else GOLD,True,align=PP_ALIGN.CENTER);text(s,label,x,4.3,3.8,1.05,23,NAVY,True,align=PP_ALIGN.CENTER)
        text(s,'Olay sayısı ≠ kişi sayısı · evaluation-only association',1.1,5.85,11.2,.4,19,MUTED,align=PP_ALIGN.CENTER)
    elif layout=='formula':
        text(s,d['formula'],.85,1.7,11.7,1,46,NAVY,True,align=PP_ALIGN.CENTER)
        text(s,d['example'],.85,3.2,11.7,.8,28,TEAL,True,align=PP_ALIGN.CENTER)
        for i,(var,label) in enumerate(d['legend']):
            x=.9+i*4.1;circle(s,x,4.8,.65,TEAL);text(s,var,x+.14,4.9,.45,.43,24,WHITE,True);text(s,label,x+.85,4.83,3.1,1.1,18)
    elif layout=='zoom':
        image(s,d['image'],.7,1.55,8.2,4.8);image(s,d['zoom'],9.4,1.85,2.85,2.85)
        rect(s,9.35,1.8,2.95,.045,GOLD);rect(s,9.35,4.7,2.95,.045,GOLD)
        text(s,d['zoom_title'],9.1,5,3.55,.45,21,NAVY,True);text(s,d['zoom_caption'],9.1,5.6,3.55,.66,17)
        if d.get('bbox'):
            iw,ih=Image.open(ASSETS/d['image']).size;k=min(8.2/iw,4.8/ih);xx=.7+(8.2-iw*k)/2;yy=1.55+(4.8-ih*k)/2
            bx,by,bw,bh=d['bbox'];line(s,xx+(bx+bw)*k,yy+(by+bh/2)*k,9.35,3.2,GOLD,1.4)
    elif layout in ('photo','diagram'):
        rect(s,.65,1.55,8.3,4.75,WHITE,radius=True)
        if layout=='photo':image(s,d['image'],.72,1.61,8.16,4.61)
        else:diagram(s,d['diagram'])
        h=min(1.37,4.72/max(1,len(items)))
        for j,(title,body) in enumerate(items):card(s,title,body,9.18,1.57+j*(h+.14),3.47,h)
    elif layout=='full':
        image(s,d['image'],.65,1.5,12,4.85)
    elif layout=='compare':
        for i,name in enumerate(d['images']):
            x=.65+i*6.15;rect(s,x,1.55,5.88,4.8,WHITE,radius=True);image(s,name,x+.08,1.65,5.72,3.6)
            title,body=items[i];text(s,title,x+.15,5.35,5.55,.38,20,NAVY,True);text(s,body,x+.15,5.81,5.55,.38,16)
    elif layout=='cards':
        count=len(items)
        if count<=3:
            w=(12.02-(count-1)*.25)/count
            for j,(title,body) in enumerate(items):
                x=.65+j*(w+.25);rect(s,x,1.8,w,3.9,WHITE,radius=True);circle(s,x+.22,2.1,.56,TEAL);text(s,str(j+1),x+.32,2.17,.36,.36,20,WHITE,True)
                text(s,title,x+.22,3.02,w-.44,.8,24,NAVY,True);text(s,body,x+.22,4.05,w-.44,1.5,21)
        else:
            for j,(title,body) in enumerate(items):card(s,title,body,.65+(j%2)*6.15,1.7+(j//2)*2.22,5.86,1.98)
    elif layout=='flow':
        count=len(items);w=(12-(count-1)*.38)/count
        for j,(title,body) in enumerate(items):
            x=.65+j*(w+.38);rect(s,x,2.05,w,3.5,WHITE,radius=True);text(s,f'{j+1:02}',x+.18,2.3,w-.36,.7,32,TEAL,True);text(s,title,x+.18,3.25,w-.36,.8,22,NAVY,True);text(s,body,x+.18,4.2,w-.36,1.1,18)
            if j<count-1:arrow(s,x+w+.03,3.6,.3,.35)
    elif layout=='table':draw_table(s,d['headers'],d['rows'])
    elif layout=='code':
        rect(s,.7,1.7,11.95,2.8,NAVY,radius=True);text(s,d['code'].replace('\n  ',' '),.98,2,11.4,2.25,17,WHITE,font='Consolas')
        for i,(title,body) in enumerate(items):card(s,title,body,.7+i*6.12,4.85,5.83,1.25)
    else:raise ValueError(layout)
    notes='ANA FİKİR: '+d['message']+'\n\n'+d['notes']
    if d.get('question'):notes+='\n\nSINIFA SORU / KONTROL: '+d['question']
    if d.get('sources'):notes+='\n\nKAYNAKLAR:\n'+'\n'.join(SOURCES[k] for k in d['sources'])
    if n in ILLUSTRATED:
        notes+='\n\nGÖRSEL: Built-in image_gen ile üretilen kavramsal öğretim çizimi; gerçek simülasyon ekran görüntüsü, ölçüm veya gerçek model renderı değildir. Dosya: sunum/uretilen_gorseller/'+ILLUSTRATED[n]['file']+'. Tam prompt ve köken: sunum/uretilen_gorseller/prompts.json ve provenance.json.'
    else:
        notes+='\n\nGÖRSEL: Gerçek Gazebo RGB kaydı, gerçek uçuş verisinden grafik veya öğretim şeması. Köken listeleri: sunum/gorseller/provenance.json, explainers-provenance.json ve gaussian-provenance.json.'
    s.notes_slide.notes_text_frame.text=notes
    return notes
notes=[]
for n,d in enumerate(SLIDES,1):notes.append(f"{'='*72}\n{n:02d}. {d['title']}\n{'='*72}\n"+render(d,n))
P.save(ROOT/'Hafta-02-Sunum.pptx')
(ROOT/'Hafta-02-Egitmen-Notlari.txt').write_text('\n\n'.join(notes),encoding='utf-8')
(ROOT/'sunum/kaynak/slides.json').write_text(json.dumps({'count':len(SLIDES),'slides':SLIDES,'sources':SOURCES},indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps({'slides':len(SLIDES),'pptx':str(ROOT/'Hafta-02-Sunum.pptx')},ensure_ascii=False))
